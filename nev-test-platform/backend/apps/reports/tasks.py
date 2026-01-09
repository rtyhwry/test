"""Celery tasks for report generation."""
from celery import shared_task
from django.conf import settings
from django.utils import timezone
from jinja2 import Environment, FileSystemLoader
import json
import os
import logging

logger = logging.getLogger(__name__)


@shared_task
def generate_report(execution_id: int, formats: list = None, 
                   include_logs: bool = True, include_screenshots: bool = True):
    """Generate test report for execution."""
    from .models import TestReport
    from apps.tasks.models import TaskExecution
    from apps.testcases.models import TestCaseResult
    
    formats = formats or ['html']
    
    try:
        execution = TaskExecution.objects.select_related(
            'task', 'environment', 'environment__host'
        ).get(id=execution_id)
        
        logger.info(f"Generating report for execution {execution_id}")
        
        # Get or create report
        report, created = TestReport.objects.get_or_create(
            execution=execution,
            defaults={
                'title': f'{execution.task.name} - 执行报告 #{execution.execution_number}'
            }
        )
        
        # Get case results
        case_results = TestCaseResult.objects.filter(
            execution=execution
        ).select_related('test_case')
        
        # Calculate statistics
        report.total_cases = case_results.count() or execution.total_cases
        report.passed_cases = case_results.filter(status='passed').count() or execution.passed_cases
        report.failed_cases = case_results.filter(status='failed').count() or execution.failed_cases
        report.skipped_cases = case_results.filter(status='skipped').count() or execution.skipped_cases
        report.blocked_cases = case_results.filter(status='blocked').count()
        report.error_cases = case_results.filter(status='error').count()
        
        if report.total_cases > 0:
            report.pass_rate = round(report.passed_cases / report.total_cases * 100, 2)
        
        report.total_duration_seconds = execution.duration_seconds
        
        # Environment info
        if execution.environment:
            report.environment_name = execution.environment.name
            report.host_info = {
                'name': execution.environment.host.name,
                'ip': execution.environment.host.ip_address,
                'os': execution.environment.host.os_type
            }
            report.device_info = [
                {
                    'name': d.name,
                    'type': d.device_type,
                    'version': d.software_version
                }
                for d in execution.environment.devices.all()
            ]
        
        # Version info
        report.software_version = execution.artifact_version
        report.git_branch = execution.task.git_branch
        report.git_commit = execution.git_commit
        
        # Generate summary
        report.summary = f"""
执行任务: {execution.task.name}
执行环境: {report.environment_name}
执行时间: {execution.start_time} - {execution.end_time}
执行状态: {execution.get_status_display()}
用例统计: 总共 {report.total_cases} 个，通过 {report.passed_cases} 个，失败 {report.failed_cases} 个
通过率: {report.pass_rate}%
        """.strip()
        
        # Ensure report directory exists
        report_dir = settings.REPORT_DIR / f'execution_{execution_id}'
        report_dir.mkdir(parents=True, exist_ok=True)
        
        # Generate reports in requested formats
        report_data = {
            'execution': execution,
            'report': report,
            'case_results': case_results,
            'generated_at': timezone.now()
        }
        
        if 'html' in formats:
            html_path = generate_html_report(report_data, report_dir)
            report.html_report_path = str(html_path)
        
        if 'json' in formats:
            json_path = generate_json_report(report_data, report_dir)
            report.json_report_path = str(json_path)
        
        if 'excel' in formats:
            excel_path = generate_excel_report(report_data, report_dir)
            report.excel_report_path = str(excel_path)
        
        report.save()
        
        logger.info(f"Report generated successfully for execution {execution_id}")
        
        return {
            'success': True,
            'report_id': report.id,
            'html_path': report.html_report_path,
            'json_path': report.json_report_path,
            'excel_path': report.excel_report_path
        }
        
    except TaskExecution.DoesNotExist:
        logger.error(f"Execution {execution_id} not found")
        return {'success': False, 'error': 'Execution not found'}
    except Exception as e:
        logger.error(f"Error generating report: {e}")
        return {'success': False, 'error': str(e)}


def generate_html_report(report_data: dict, output_dir) -> str:
    """Generate HTML report."""
    html_content = f"""
<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{report_data['report'].title}</title>
    <style>
        body {{ font-family: 'Microsoft YaHei', Arial, sans-serif; margin: 0; padding: 20px; background: #f5f5f5; }}
        .container {{ max-width: 1200px; margin: 0 auto; background: white; padding: 30px; border-radius: 8px; box-shadow: 0 2px 10px rgba(0,0,0,0.1); }}
        h1 {{ color: #333; border-bottom: 2px solid #007bff; padding-bottom: 10px; }}
        h2 {{ color: #555; margin-top: 30px; }}
        .summary {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(200px, 1fr)); gap: 20px; margin: 20px 0; }}
        .stat-card {{ background: #f8f9fa; padding: 20px; border-radius: 8px; text-align: center; }}
        .stat-card .value {{ font-size: 36px; font-weight: bold; color: #007bff; }}
        .stat-card .label {{ color: #666; margin-top: 5px; }}
        .stat-card.success .value {{ color: #28a745; }}
        .stat-card.danger .value {{ color: #dc3545; }}
        .pass-rate {{ font-size: 48px; font-weight: bold; }}
        .pass-rate.high {{ color: #28a745; }}
        .pass-rate.medium {{ color: #ffc107; }}
        .pass-rate.low {{ color: #dc3545; }}
        table {{ width: 100%; border-collapse: collapse; margin-top: 20px; }}
        th, td {{ padding: 12px; text-align: left; border-bottom: 1px solid #ddd; }}
        th {{ background: #007bff; color: white; }}
        tr:hover {{ background: #f5f5f5; }}
        .status-passed {{ color: #28a745; font-weight: bold; }}
        .status-failed {{ color: #dc3545; font-weight: bold; }}
        .status-skipped {{ color: #6c757d; }}
        .info-section {{ background: #f8f9fa; padding: 15px; border-radius: 8px; margin: 20px 0; }}
        .footer {{ text-align: center; color: #999; margin-top: 30px; padding-top: 20px; border-top: 1px solid #eee; }}
    </style>
</head>
<body>
    <div class="container">
        <h1>📊 {report_data['report'].title}</h1>
        
        <div class="info-section">
            <p><strong>执行时间:</strong> {report_data['execution'].start_time} - {report_data['execution'].end_time}</p>
            <p><strong>执行环境:</strong> {report_data['report'].environment_name}</p>
            <p><strong>软件版本:</strong> {report_data['report'].software_version or 'N/A'}</p>
            <p><strong>Git分支:</strong> {report_data['report'].git_branch} ({report_data['report'].git_commit[:8] if report_data['report'].git_commit else 'N/A'})</p>
        </div>
        
        <h2>📈 执行统计</h2>
        <div class="summary">
            <div class="stat-card">
                <div class="value">{report_data['report'].total_cases}</div>
                <div class="label">总用例数</div>
            </div>
            <div class="stat-card success">
                <div class="value">{report_data['report'].passed_cases}</div>
                <div class="label">通过</div>
            </div>
            <div class="stat-card danger">
                <div class="value">{report_data['report'].failed_cases}</div>
                <div class="label">失败</div>
            </div>
            <div class="stat-card">
                <div class="value">{report_data['report'].skipped_cases}</div>
                <div class="label">跳过</div>
            </div>
        </div>
        
        <div style="text-align: center; margin: 30px 0;">
            <div class="pass-rate {'high' if report_data['report'].pass_rate >= 80 else 'medium' if report_data['report'].pass_rate >= 60 else 'low'}">
                {report_data['report'].pass_rate}%
            </div>
            <div style="color: #666;">通过率</div>
        </div>
        
        <h2>📋 用例详情</h2>
        <table>
            <thead>
                <tr>
                    <th>用例编号</th>
                    <th>用例名称</th>
                    <th>状态</th>
                    <th>执行时长</th>
                </tr>
            </thead>
            <tbody>
"""
    
    for result in report_data['case_results']:
        status_class = f"status-{result.status}"
        duration = f"{result.duration_ms}ms" if result.duration_ms else "N/A"
        html_content += f"""
                <tr>
                    <td>{result.test_case.case_code}</td>
                    <td>{result.test_case.name}</td>
                    <td class="{status_class}">{result.get_status_display()}</td>
                    <td>{duration}</td>
                </tr>
"""
    
    html_content += f"""
            </tbody>
        </table>
        
        <div class="footer">
            <p>生成时间: {report_data['generated_at'].strftime('%Y-%m-%d %H:%M:%S')}</p>
            <p>新能源汽车测试平台</p>
        </div>
    </div>
</body>
</html>
"""
    
    output_path = output_dir / 'report.html'
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(html_content)
    
    return str(output_path)


def generate_json_report(report_data: dict, output_dir) -> str:
    """Generate JSON report."""
    report = report_data['report']
    execution = report_data['execution']
    
    json_data = {
        'title': report.title,
        'execution': {
            'id': execution.id,
            'task_name': execution.task.name,
            'execution_number': execution.execution_number,
            'status': execution.status,
            'start_time': str(execution.start_time) if execution.start_time else None,
            'end_time': str(execution.end_time) if execution.end_time else None,
            'duration_seconds': execution.duration_seconds
        },
        'statistics': {
            'total_cases': report.total_cases,
            'passed_cases': report.passed_cases,
            'failed_cases': report.failed_cases,
            'skipped_cases': report.skipped_cases,
            'blocked_cases': report.blocked_cases,
            'error_cases': report.error_cases,
            'pass_rate': report.pass_rate
        },
        'environment': {
            'name': report.environment_name,
            'host': report.host_info,
            'devices': report.device_info
        },
        'version': {
            'software': report.software_version,
            'git_branch': report.git_branch,
            'git_commit': report.git_commit
        },
        'case_results': [
            {
                'case_code': r.test_case.case_code,
                'case_name': r.test_case.name,
                'status': r.status,
                'duration_ms': r.duration_ms,
                'error_message': r.error_message
            }
            for r in report_data['case_results']
        ],
        'generated_at': str(report_data['generated_at'])
    }
    
    output_path = output_dir / 'report.json'
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(json_data, f, ensure_ascii=False, indent=2)
    
    return str(output_path)


def generate_excel_report(report_data: dict, output_dir) -> str:
    """Generate Excel report."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    
    wb = Workbook()
    ws = wb.active
    ws.title = "测试报告"
    
    # Header style
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="007BFF", end_color="007BFF", fill_type="solid")
    
    # Summary section
    ws['A1'] = "测试报告摘要"
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:D1')
    
    summary_data = [
        ['任务名称', report_data['execution'].task.name],
        ['执行环境', report_data['report'].environment_name],
        ['总用例数', report_data['report'].total_cases],
        ['通过数', report_data['report'].passed_cases],
        ['失败数', report_data['report'].failed_cases],
        ['通过率', f"{report_data['report'].pass_rate}%"],
    ]
    
    for i, row in enumerate(summary_data, start=3):
        ws.cell(row=i, column=1, value=row[0]).font = Font(bold=True)
        ws.cell(row=i, column=2, value=row[1])
    
    # Case results section
    start_row = len(summary_data) + 5
    ws.cell(row=start_row, column=1, value="用例执行详情").font = Font(bold=True, size=14)
    ws.merge_cells(f'A{start_row}:D{start_row}')
    
    headers = ['用例编号', '用例名称', '执行状态', '执行时长(ms)']
    header_row = start_row + 2
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    
    for i, result in enumerate(report_data['case_results'], start=header_row + 1):
        ws.cell(row=i, column=1, value=result.test_case.case_code)
        ws.cell(row=i, column=2, value=result.test_case.name)
        ws.cell(row=i, column=3, value=result.get_status_display())
        ws.cell(row=i, column=4, value=result.duration_ms or 'N/A')
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    
    output_path = output_dir / 'report.xlsx'
    wb.save(output_path)
    
    return str(output_path)


@shared_task
def sync_report_to_alm(report_id: int):
    """Sync report results to ALM system."""
    from .models import TestReport
    
    try:
        report = TestReport.objects.select_related('execution', 'execution__task').get(id=report_id)
        
        # This would integrate with actual ALM API
        # For now, just mark as synced
        logger.info(f"Syncing report {report_id} to ALM")
        
        # TODO: Implement actual ALM sync logic
        # alm_client = ALMClient(settings.ALM_URL, settings.ALM_USERNAME, settings.ALM_PASSWORD)
        # alm_client.update_test_run(report.execution.task.alm_test_set_id, results)
        
        report.alm_synced = True
        report.alm_sync_time = timezone.now()
        report.save()
        
        return {'success': True, 'report_id': report_id}
        
    except TestReport.DoesNotExist:
        return {'success': False, 'error': 'Report not found'}
    except Exception as e:
        logger.error(f"Error syncing report to ALM: {e}")
        return {'success': False, 'error': str(e)}
